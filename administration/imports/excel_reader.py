from openpyxl import load_workbook
from django import forms


class ExcelReader(forms.Form):
    
    excel_file = forms.FileField(label="Excel-Datei auswählen")

    '''
    def __init__(self, file):
        self.file = file
    '''

    def read(self):
        
        uploaded_file = self.cleaned_data['excel_file']

        COLUMN_DATUM_TTB = 0
        COLUMN_EG_ART = 2
        COLUMN_MANUFACTURER = 3
        COLUMN_SERIALNUMBER = 5
        COLUMN_TEI = 6
        COLUMN_ISSI = 9
        COLUMN_SIKANUMBER = 10
        COLUMN_SIKANAME = 11
        COLUMN_UNIT = 12
        COLUMN_FUNKRUFNAME = 13
        COLUMN_USECASE = 14
        COLUMN_LAND_BY = 16
        COLUMN_ORG_FW = 17
        COLUMN_LANDKREIS = 18
        

        workbook = load_workbook(
            uploaded_file,
            data_only=True,
        )

        sheet = workbook.active
        rows = []

        # Gerätetabelle beginnt bei Zeile 5
        for excel_row in sheet.iter_rows(min_row=5, values_only=True):
            
            funkrufname_komplett = str(excel_row[COLUMN_FUNKRUFNAME]).strip() if excel_row[COLUMN_FUNKRUFNAME] is not None else ""
            
            ort_name = ""
            
            if funkrufname_komplett:
                
                partition = funkrufname_komplett.split(" ")
                
                if len(partition) >= 2:
                    ort_name = partition[1]
                    
                else:
                    ort_name = partition[0]
                    
            row = {                
                "datum_ttb": str(excel_row[COLUMN_DATUM_TTB]).strip() if excel_row[COLUMN_DATUM_TTB] else "",
                "eg_art": str(excel_row[COLUMN_EG_ART]).strip() if excel_row[COLUMN_EG_ART] else "",
                "hersteller": str(excel_row[COLUMN_MANUFACTURER]).strip() if excel_row[COLUMN_MANUFACTURER] else "",
                "seriennummer": str(excel_row[COLUMN_SERIALNUMBER]).strip() if excel_row[COLUMN_SERIALNUMBER] else "",
                "tei": str(excel_row[COLUMN_TEI]).strip() if excel_row[COLUMN_TEI] else "",
                "issi": str(excel_row[COLUMN_ISSI]).strip() if excel_row[COLUMN_ISSI] else "",
                "sika-nummer": str(excel_row[COLUMN_SIKANUMBER]).strip() if excel_row[COLUMN_SIKANUMBER] else "",
                "sika-name": str(excel_row[COLUMN_SIKANAME]).strip() if excel_row[COLUMN_SIKANAME] else "",
                "fahrzeugart": str(excel_row[COLUMN_UNIT]).strip() if excel_row[COLUMN_UNIT] else "",
                "kommune": ort_name,
                "funkrufname": funkrufname_komplett,
                "verwendung": str(excel_row[COLUMN_USECASE]).strip() if excel_row[COLUMN_USECASE] else "",
                "land": str(excel_row[COLUMN_LAND_BY]).strip() if excel_row[COLUMN_LAND_BY] else "",
                "org": str(excel_row[COLUMN_ORG_FW]).strip() if excel_row[COLUMN_ORG_FW] else "",
                "landkreis": str(excel_row[COLUMN_LANDKREIS]).strip() if excel_row[COLUMN_LANDKREIS] else "",
                
                "gopta_itsi": str(excel_row[15]).strip() if excel_row[15] else "",

                "aopta_t": str(excel_row[19]).strip() if excel_row[19] else "",
                "aopta_u": str(excel_row[20]).strip() if excel_row[20] else "",
                "aopta_v": str(excel_row[21]).strip() if excel_row[21] else "",
                "aopta_w": str(excel_row[22]).strip() if excel_row[22] else "",
                "aopta_x": str(excel_row[23]).strip() if excel_row[23] else "",
                "aopta_y": str(excel_row[24]).strip() if excel_row[24] else "",
                "aopta_z": str(excel_row[25]).strip() if excel_row[25] else "",
                "aopta_aa": str(excel_row[26]).strip() if excel_row[26] else "",
                "aopta_ab": str(excel_row[27]).strip() if excel_row[27] else "",
                "aopta_ac": str(excel_row[28]).strip() if excel_row[28] else "",
                "aopta_ad": str(excel_row[29]).strip() if excel_row[29] else "",
                "aopta_ae": str(excel_row[30]).strip() if excel_row[30] else "",
                "aopta_af": str(excel_row[31]).strip() if excel_row[31] else "",
                "aopta_ag": str(excel_row[32]).strip() if excel_row[32] else "",
                "aopta_ah": str(excel_row[33]).strip() if excel_row[33] else "",
                "aopta_ai": str(excel_row[34]).strip() if excel_row[33] else "",
            }

            if not excel_row[5]:
                break

            rows.append(row)

        return rows