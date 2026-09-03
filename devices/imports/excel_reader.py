from openpyxl import load_workbook
from enum import Enum

class ImportType(Enum):
    ENDGERAETE = "endgeraete"
    SIRENEN = "sirenen"


class ExcelReader:

    def __init__(self, file):
        self.file = file

    def read(self):

        COLUMN_EG_ART = 1
        COLUMN_MANUFACTURER = 2
        COLUMN_DEVICE_TYPE = 3
        COLUMN_SERIALNUMBER = 4
        COLUMN_TEI = 5
        COLUMN_ISSI = 8
        COLUMN_SIKANUMBER = 9
        COLUMN_SIKANAME = 10
        COLUMN_UNIT = 11
        COLUMN_FUNKRUFNAME = 12
        COLUMN_USECASE = 13
        

        workbook = load_workbook(
            self.file,
            data_only=True,
        )

        sheet = workbook.active

        # Stammdaten aus dem Formular
        landkreis = sheet["D4"].value
        kommune = sheet["D6"].value
        organisationsart = sheet["D7"].value
        organisationsname = sheet["D8"].value

        rows = []

        # Gerätetabelle beginnt bei Zeile 18
        title = str(sheet["A2"].value or "")
        if "Endgeräte" in title:
            for excel_row in sheet.iter_rows(min_row=18, values_only=True):
                opta = "".join(
                    str(excel_row[i]).strip()
                    for i in range(18, 23)
                    if excel_row[i]
                )
                opta2 = "".join(
                    str(excel_row[i]).strip()
                    for i in range(24, 34)
                    if excel_row[i]
                )
                row = {
                    "landkreis": str(landkreis).strip() if landkreis else "",
                    "kommune": str(kommune).strip() if kommune else "",
                    "organisationsart": str(organisationsart).strip() if organisationsart else "",
                    "organisationsname": str(organisationsname).strip() if organisationsname else "",
                    
                    "eg_art": str(excel_row[COLUMN_EG_ART]).strip() if excel_row[COLUMN_EG_ART] else "",
                    "hersteller": str(excel_row[COLUMN_MANUFACTURER]).strip() if excel_row[COLUMN_MANUFACTURER] else "",

                    # Diese Spalten müssen wir ggf. noch anpassen,
                    # sobald wir die exakten Spalten kennen.
                    "seriennummer": str(excel_row[COLUMN_SERIALNUMBER]).strip() if excel_row[COLUMN_SERIALNUMBER] else "",
                    "tei": str(excel_row[COLUMN_TEI]).strip() if excel_row[COLUMN_TEI] else "",
                    "opta": opta,
                    "issi": str(excel_row[COLUMN_ISSI]).strip() if excel_row[COLUMN_ISSI] else "",
                    "sika-nummer": str(excel_row[COLUMN_SIKANUMBER]).strip() if excel_row[COLUMN_SIKANUMBER] else "",
                    "sika-name": str(excel_row[COLUMN_SIKANAME]).strip() if excel_row[COLUMN_SIKANAME] else "",
                    "fahrzeugart": str(excel_row[COLUMN_UNIT]).strip() if excel_row[COLUMN_UNIT] else "",
                    "funkrufname": str(excel_row[COLUMN_FUNKRUFNAME]).strip() if excel_row[COLUMN_FUNKRUFNAME] else "",
                    "verwendung": str(excel_row[COLUMN_USECASE]).strip() if excel_row[COLUMN_USECASE] else "",
                    
                    "gopta_itsi": str(excel_row[14]).strip() if excel_row[14] else "",
                    "aopta_land": str(excel_row[15]).strip() if excel_row[15] else "",
                    "aopta_org": str(excel_row[16]).strip() if excel_row[16] else "",
                    "aopta_region": str(excel_row[17]).strip() if excel_row[17] else "",

                    "aopta_t": str(excel_row[18]).strip() if excel_row[18] else "",
                    "aopta_u": str(excel_row[19]).strip() if excel_row[19] else "",
                    "aopta_v": str(excel_row[20]).strip() if excel_row[20] else "",
                    "aopta_w": str(excel_row[21]).strip() if excel_row[21] else "",
                    "aopta_x": str(excel_row[22]).strip() if excel_row[22] else "",
                    "aopta_y": str(excel_row[23]).strip() if excel_row[23] else "",
                    "aopta_z": str(excel_row[24]).strip() if excel_row[24] else "",
                    "aopta_aa": str(excel_row[25]).strip() if excel_row[25] else "",
                    "aopta_ab": str(excel_row[26]).strip() if excel_row[26] else "",
                    "aopta_ac": str(excel_row[27]).strip() if excel_row[27] else "",
                    "aopta_ad": str(excel_row[28]).strip() if excel_row[28] else "",
                    "aopta_ae": str(excel_row[29]).strip() if excel_row[29] else "",
                    "aopta_af": str(excel_row[30]).strip() if excel_row[30] else "",
                    "aopta_ag": str(excel_row[31]).strip() if excel_row[31] else "",
                    "aopta_ah": str(excel_row[32]).strip() if excel_row[32] else "",
                    "aopta_ai": str(excel_row[33]).strip() if excel_row[33] else "",
                }
                print(row)
                if not excel_row[5]:
                    break

                rows.append(row)

            return rows, ImportType.ENDGERAETE
    
        elif "Sirenen" in title:
            
            return rows, ImportType.SIRENEN
        
        else:
            raise ValueError(
                    f"Unbekannter Antragstyp: {title}"
                )