from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Color
from openpyxl.utils import get_column_letter
from devices.constants import DeviceConstants as Constants
from devices.models import Device, ImportStatus
from pathlib import Path
from copy import copy


class RadioDirectoryExporter:

    MRT = "MRT"
    FRT = "FRT"

    DOWNLOAD_FOLDER = Path.home() / "Downloads"

    def __init__(self, landkreis):

        self.landkreis = landkreis

        status = ImportStatus.objects.get(
            import_type="endgeraete"
        )

        self.import_date = status.last_import.replace(
            tzinfo=None
        )

    def export(self):

        devices = Device.objects.filter(
            landkreis=self.landkreis
        )

        self.DOWNLOAD_FOLDER.mkdir(
            parents=True,
            exist_ok=True,
        )

        filename = (
            f"{self.import_date:%Y-%m-%d}"
            f"_Funkverzeichnis_Digitalfunk_"
            f"{self.landkreis}.xlsx"
        )

        filepath = self.DOWNLOAD_FOLDER / filename

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Datenimport"

        self._create_header(sheet)
        self._write_devices(sheet, devices)
        self._format_sheet(sheet)

        workbook.save(filepath)

        return filepath

    # ---------------------------------------------------------
    # HEADER
    # ---------------------------------------------------------

    def _create_header(self, sheet):

        headers = {
            "A1": "Datum        TTB",
            "B1": "Datum        ILS",
            "C1": "EG-Art",
            "D1": "Hersteller",
            "E1": "EG-Typ",
            "F1": "Seriennummer",
            "G1": "TEI",
            "H1": "MCC",
            "I1": "MNC",
            "J1": "SSI",
            "K1": "BOS-SiKa-Nummer",
            "L1": "BOS-SiKa-Name",
            "M1": "Fahrzeugart/Standort",
            "N1": "Funkrufname des Fahrzeugs/Standorts     max 32 Zeichen!",
            "O1": "Verwendung",
            "P1": "GOPTA-ITSI",
            "Q1": "AOPTA-Land",
            "R1": "AOPTA-Org",
            "S1": "AOPTA-Region",
            "T1": "AOPTA-Zusatz",
            "AJ1": "Bemerkung",
        }

        for cell, value in headers.items():
            sheet[cell] = value

        # A bis S über vier Zeilen
        for column in range(1, 20):
            letter = get_column_letter(column)

            sheet.merge_cells(
                f"{letter}1:{letter}4"
            )

        # AOPTA-Zusatz
        sheet.merge_cells("T1:AI1")

        # AOPTA-Zusatz dritte Ebene
        sheet.merge_cells("T3:AI3")

        # Bemerkung
        sheet.merge_cells("AJ1:AJ4")

        # AOPTA-Zusatz zweite Ebene
        for column, value in enumerate(
            range(9, 25),
            start=20,
        ):
            sheet.cell(
                row=2,
                column=column,
                value=value,
            )

        # AOPTA-Zusatz vierte Ebene
        sheet.merge_cells("T4:X4")
        sheet["T4"] = "4.1"

        sheet.merge_cells("Y4:AF4")
        sheet["Y4"] = "4.2"

        sheet.merge_cells("AG4:AH4")
        sheet["AG4"] = "4.3"

        sheet["AI4"] = 5

    # ---------------------------------------------------------
    # GERÄTEDATEN
    # ---------------------------------------------------------

    def _write_devices(self, sheet, devices):

        start_row = 5

        for row_number, device in enumerate(
            devices,
            start=start_row,
        ):

            sheet[f"A{row_number}"] = (
                device.datum_ttb.replace(tzinfo=None)
                if device.datum_ttb
                else None
            )

            sheet[f"B{row_number}"] = self._get(
                device,
                "datum_ils",
            )

            sheet[f"C{row_number}"] = self._get(
                device,
                "eg_art",
            )

            sheet[f"D{row_number}"] = self._get(
                device,
                "hersteller",
            )

            sheet[f"E{row_number}"] = self._get(
                device,
                "eg_typ",
            )

            sheet[f"F{row_number}"] = self._get(
                device,
                "seriennummer",
            )

            sheet[f"G{row_number}"] = self._get(
                device,
                "tei",
            )

            # Globale Konstanten
            sheet[f"H{row_number}"] = Constants.MCC
            sheet[f"I{row_number}"] = Constants.MNC

            sheet[f"J{row_number}"] = self._get(
                device,
                "issi",
            )

            sheet[f"K{row_number}"] = self._get(
                device,
                "bos_sika_nummer",
            )

            sheet[f"L{row_number}"] = self._get(
                device,
                "bos_sika_name",
            )

            sheet[f"M{row_number}"] = self._get(
                device,
                "fahrzeugart",
            )

            sheet[f"N{row_number}"] = self._get(
                device,
                "funkrufname",
            )

            sheet[f"O{row_number}"] = self._get(
                device,
                "verwendung",
            )

            sheet[f"P{row_number}"] = self._get(
                device,
                "gopta_itsi",
            )

            sheet[f"Q{row_number}"] = self._get(
                device,
                "aopta_land",
            )

            sheet[f"R{row_number}"] = self._get(
                device,
                "aopta_org",
            )

            sheet[f"S{row_number}"] = self._get(
                device,
                "aopta_region",
            )

            # AOPTA-Zusatz
            for column in range(20, 36):

                letter = get_column_letter(column)

                value = self._get(
                    device,
                    f"aopta_{letter.lower()}",
                )

                sheet[
                    f"{letter}{row_number}"
                ] = value

            # Bemerkung
            sheet[f"AJ{row_number}"] = self._get(
                device,
                "bemerkung",
            )

            self._apply_row_style(
                sheet,
                row_number,
                self._get(
                    device,
                    "eg_art",
                ),
            )

    # ---------------------------------------------------------
    # ZEILENSTYLES
    # ---------------------------------------------------------

    def _apply_row_style(
        self,
        sheet,
        row_number,
        eg_art,
    ):

        if eg_art == self.MRT:

            self._apply_mrt_style(
                sheet,
                row_number,
            )

        elif eg_art == self.FRT:

            self._apply_frt_style(
                sheet,
                row_number,
            )

    def _apply_mrt_style(self, sheet, row_number):
        fill = PatternFill(
            fill_type="solid",
            fgColor=Color(theme=7, tint=0.5999938962981048),
        )
        for column in range(3, 16):
            sheet.cell(row=row_number, column=column).fill = copy(fill)


    def _apply_frt_style(self, sheet, row_number):
        fill = PatternFill(
            fill_type="solid",
            fgColor=Color(theme=9, tint=0.3999755851924192),
        )
        for column in range(3, 16):
            sheet.cell(row=row_number, column=column).fill = copy(fill)

    # ---------------------------------------------------------
    # FORMATIERUNG
    # ---------------------------------------------------------

    def _format_sheet(self, sheet):
        """
        Übernimmt das Grundlayout des ursprünglichen Funkverzeichnisses.
        Dynamische Farben für MRT/FRT werden separat gesetzt.
        """

        # ---------------------------------------------------------
        # Schrift und Grundausrichtung
        # ---------------------------------------------------------

        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=36):
            for cell in row:
                cell.font = Font(name="Calibri", size=11)

        for row in range(5, sheet.max_row + 1):

            # C:J zentriert
            for column in range(3, 11):
                sheet.cell(row=row, column=column).alignment = Alignment(
                    horizontal="center"
                )

            # N linksbündig
            sheet.cell(row=row, column=14).alignment = Alignment(
                horizontal="left"
            )

        # ---------------------------------------------------------
        # Kopfzeile
        # ---------------------------------------------------------

        for row in sheet.iter_rows(
            min_row=1,
            max_row=4,
            min_col=1,
            max_col=36,
        ):
            for cell in row:
                cell.font = Font(
                    name="Calibri",
                    size=11,
                    bold=True,
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

        header_fill = PatternFill(
            fill_type="solid",
            fgColor=Color(
                theme=0,
                tint=-0.3499862666707358,
            ),
        )

        for row in sheet.iter_rows(
            min_row=1,
            max_row=4,
            min_col=1,
            max_col=19,  # A:S
        ):
            for cell in row:
                cell.fill = copy(header_fill)

        for row in sheet.iter_rows(
            min_row=1,
            max_row=4,
            min_col=36,
            max_col=36,  # AJ
        ):
            for cell in row:
                cell.fill = copy(header_fill)
        
        aopta_fill = PatternFill(
            fill_type="solid",
            fgColor="FFFFCC",
        )

        for cell in ["T1", "T2", "T3", "T4", "Y4", "AG4", "AI4"]:
            sheet[cell].fill = copy(aopta_fill)
        
        aopta_header_font = Font(
            name="Arial",
            size=10,
            bold=True,
        )

        for cell in ["T1", "T3", "T4", "Y4", "AG4", "AI4"]:
            sheet[cell].font = copy(aopta_header_font)
        
        aopta_hair = Side(style="hair", color="000000")
        aopta_thin = Side(style="thin", color="000000")

        for row in sheet.iter_rows(
            min_row=1,
            max_row=4,
            min_col=20,
            max_col=35,
        ):
            for cell in row:
                cell.border = Border(
                    left=aopta_hair,
                    right=aopta_hair,
                    top=aopta_hair,
                    bottom=aopta_hair,
                )

        # äußerer rechter Rand von AI
        for row in range(1, 5):
            sheet.cell(row=row, column=35).border = Border(
                left=aopta_hair,
                right=aopta_thin,
                top=aopta_hair,
                bottom=aopta_hair,
            )
        
        # AOPTA-Ausrichtung wie im Original
        sheet["T1"].alignment = Alignment(horizontal="center", vertical="center")
        sheet["T2"].alignment = Alignment(horizontal="center")
        sheet["T3"].alignment = Alignment(horizontal="center", vertical="center")
        sheet["T4"].alignment = Alignment(horizontal="center", vertical="center")
        sheet["Y4"].alignment = Alignment(horizontal="center", vertical="center")
        sheet["AG4"].alignment = Alignment(horizontal="center", vertical="center")
        sheet["AI4"].alignment = Alignment(horizontal="center")
        
        for cell in sheet[2][19:35]:
            cell.fill = copy(aopta_fill)
        
        for cell in sheet[2][19:35]:
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(horizontal="center")
        
        # ---------------------------------------------------------
        # Rahmen
        # ---------------------------------------------------------

        thin_side = Side(
            style="thin",
            color="000000",
        )

        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for row in sheet.iter_rows(
            min_row=1,
            max_row=4,
            min_col=1,
            max_col=19
        ):
            for cell in row:
                cell.border = border

        for row in sheet.iter_rows(
            min_row=1,
            max_row=4,
            min_col=36,
            max_col=36
        ):
            for cell in row:
                cell.border = border

        # ---------------------------------------------------------
        # Datum
        # ---------------------------------------------------------

        for row in range(5, sheet.max_row + 1):
            sheet[f"A{row}"].number_format = "DD.MM.YYYY"
            sheet[f"B{row}"].number_format = "DD.MM.YYYY"

        # ---------------------------------------------------------
        # Spaltenbreiten
        # ---------------------------------------------------------

        widths = {
            "A": 12.1328125,
            "B": 10.1328125,
            "C": 11.3984375,
            "D": 16.265625,
            "E": 11.86328125,
            "F": 18.86328125,
            "G": 16.86328125,
            "H": 9.73046875,
            "I": 10.0,
            "J": 13.0,
            "K": 26.1328125,
            "L": 25.86328125,
            "M": 22.86328125,
            "N": 35.0,
            "O": 24.86328125,
            "P": 25.59765625,
            "Q": 11.3984375,
            "R": 11.59765625,
            "S": 12.73046875,
            "T": 3.265625,
            "U": 3.1328125,
            "V": 13.0,
            "W": 13.0,
            "X": 13.0,
            "Y": 13.0,
            "Z": 13.0,
            "AA": 13.0,
            "AB": 13.0,
            "AC": 13.0,
            "AD": 13.0,
            "AE": 13.0,
            "AF": 13.0,
            "AG": 13.0,
            "AH": 13.0,
            "AI": 13.0,
            "AJ": 79.0,
        }

        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        # ---------------------------------------------------------
        # Zeilenhöhen
        # ---------------------------------------------------------

        sheet.row_dimensions[1].height = 44.25
        sheet.row_dimensions[2].height = 15.0
        sheet.row_dimensions[3].height = 15.0
        sheet.row_dimensions[4].height = 15.0

        # Die Datenzeilen entsprechen dem Original
        for row in range(5, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 15.0

        # ---------------------------------------------------------
        # Ansicht
        # ---------------------------------------------------------

        sheet.freeze_panes = "A5"

        # ---------------------------------------------------------
        # Drucklayout
        # ---------------------------------------------------------

        sheet.page_setup.orientation = "landscape"

        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.75
        sheet.page_margins.bottom = 0.75

        sheet.print_area = f"A1:AJ{sheet.max_row}"

        # ---------------------------------------------------------
        # Keine automatische Skalierung erzwingen
        # ---------------------------------------------------------

        sheet.page_setup.fitToWidth = None
        sheet.page_setup.fitToHeight = None

    # ---------------------------------------------------------
    # HILFSMETHODE
    # ---------------------------------------------------------

    def _get(self, device, *names):

        for name in names:

            value = getattr(
                device,
                name,
                None,
            )

            if value is not None:
                return value

        return ""
