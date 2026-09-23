from pathlib import Path
import re

from openpyxl import load_workbook
from pypdf import PdfReader


class FertigmeldungReader:

    def __init__(self, uploaded_file):
        self.uploaded_file = uploaded_file

    def read(self):
        dateiname = Path(self.uploaded_file.name)

        if dateiname.suffix.lower() == ".pdf":
            return self._read_haeusler_pdf()

        if dateiname.suffix.lower() in [".xlsx", ".xlsm"]:
            return self._read_hoermann_excel()

        raise ValueError(
            "Unbekanntes Format für Fertigmeldung."
        )

    def _read_haeusler_pdf(self):
        reader = PdfReader(self.uploaded_file)

        text = ""

        for page in reader.pages:
            page_text = page.extract_text() or ""
            text += page_text + "\n"

        tei_match = re.search(
            r"TEI\s+([0-9]+)",
            text,
        )

        datum_match = re.search(
            r"Installation\s+durchgeführt.*?(\d{2}\.\d{2}\.\d{4})",
            text,
            re.DOTALL,
        )

        if not tei_match:
            raise ValueError(
                "TEI konnte in der Häusler-PDF nicht gefunden werden."
            )

        if not datum_match:
            raise ValueError(
                "Installationsdatum konnte in der Häusler-PDF "
                "nicht gefunden werden."
            )

        return {
            "hersteller": "Haeusler",
            "tei": tei_match.group(1),
            "fertigmeldedatum": datum_match.group(1),
        }

    def _read_hoermann_excel(self):
        workbook = load_workbook(
            self.uploaded_file,
            data_only=True,
        )

        if "Tabelle1" not in workbook.sheetnames:
            raise ValueError(
                "Das erwartete Tabellenblatt 'Tabelle1' "
                "wurde nicht gefunden."
            )

        sheet = workbook["Tabelle1"]

        rows = list(
            sheet.iter_rows(
                values_only=True
            )
        )

        if not rows:
            raise ValueError(
                "Die Hörmann-Datei enthält keine Daten."
            )

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in rows[1]
        ]

        try:
            adresse_index = headers.index("Adresse")
            issi_index = headers.index("Lokale ISSI")
            datum_index = headers.index("Datum")
        except ValueError as exc:
            raise ValueError(
                "Benötigte Spalten in der Hörmann-Datei "
                "nicht gefunden."
            ) from exc

        fertigmeldungen = []

        for row in rows[2:]:

            if not any(value is not None for value in row):
                continue

            issi = row[issi_index]
            datum = row[datum_index]
            adresse = row[adresse_index]

            if issi is None:
                continue

            fertigmeldungen.append(
                {
                    "hersteller": "Hoermann",
                    "issi": str(issi).strip(),
                    "fertigmeldedatum": datum,
                    "adresse": (
                        str(adresse).strip()
                        if adresse is not None
                        else ""
                    ),
                }
            )

        if not fertigmeldungen:
            raise ValueError(
                "Keine Hörmann-Fertigmeldungen gefunden."
            )

        return fertigmeldungen