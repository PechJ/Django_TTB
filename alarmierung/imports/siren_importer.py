from alarmierung.models import SirenFRTAntrag, SirenFRTFreigabePending
from devices.models import SirenStammdaten, SirenFRTStatus
from pathlib import Path
from openpyxl import load_workbook


def erkenne_frt_dateityp(uploaded_file):
    suffix = Path(uploaded_file.name).suffix.lower()

    if suffix == ".pdf":
        return "HAEUSLER_FERTIGMELDUNG"

    if suffix not in [".xlsx", ".xlsm"]:
        return None

    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    sheetnames = workbook.sheetnames

    if "FRT_Import_Netsite_Pega" in sheetnames:
        return "FRT_ANTRAG"

    if "FRT-Daten" in sheetnames:
        return "FRT_FREIGABE"

    if "Tabelle1" in sheetnames:
        return "HOERMANN_FERTIGMELDUNG"

    return None


def normalisiere_adresse(adresse):
    if not adresse:
        return ""

    adresse = adresse.lower().strip()

    adresse = adresse.replace("straße", "strasse")
    adresse = adresse.replace("str.", "strasse")
    adresse = adresse.replace("ß", "ss")

    adresse = " ".join(adresse.split())

    return adresse


def finde_sirene(antrag):
    gesuchte_adresse = normalisiere_adresse(
        f"{antrag.strasse} {antrag.hausnummer}"
    )

    gesuchter_ort = normalisiere_adresse(
        antrag.ort
    )

    treffer = []

    for sirene in SirenStammdaten.objects.all():

        vorhandene_adresse = normalisiere_adresse(
            sirene.standort_adresse
        )

        vorhandene_gemeinde = normalisiere_adresse(
            sirene.gemeinde
        )

        vorhandener_ortsteil = normalisiere_adresse(
            sirene.ortsteil
        )

        ort_passt = (
            gesuchter_ort == vorhandene_gemeinde
            or gesuchter_ort == vorhandener_ortsteil
        )

        if (
            gesuchte_adresse == vorhandene_adresse
            and ort_passt
        ):
            treffer.append(sirene)

    return treffer


class SirenFRTAntragImporter:

    def __init__(self, rows, hersteller):
        self.rows = rows
        self.hersteller = hersteller
        self.created = 0
        self.status_created = 0
        self.waiting_created = 0

    def run(self):

        for row in self.rows:

            antrag = SirenFRTAntrag(
                suchkreisname=row["suchkreisname"],
                ort=row["ort"],
                plz=row["plz"],
                strasse=row["strasse"],
                hausnummer=row["hausnummer"],
                breitengrad_grad=row["breitengrad_grad"],
                breitengrad_min=row["breitengrad_min"],
                breitengrad_sek=row["breitengrad_sek"],
                laengengrad_grad=row["laengengrad_grad"],
                laengengrad_min=row["laengengrad_min"],
                laengengrad_sek=row["laengengrad_sek"],
                standorthoehe=row["standorthoehe"],
                antragsdatum=row["antragsdatum"],
            )

            sirenen = finde_sirene(antrag)

            if sirenen:

                sirene = sirenen[0]

                def dms_zu_dezimal(grad, minuten, sekunden):
                    if grad is None:
                        return None

                    return (
                        float(grad)
                        + float(minuten or 0) / 60
                        + float(sekunden or 0) / 3600
                    )

                sirene.breitengrad = dms_zu_dezimal(
                    antrag.breitengrad_grad,
                    antrag.breitengrad_min,
                    antrag.breitengrad_sek,
                )

                sirene.laengengrad = dms_zu_dezimal(
                    antrag.laengengrad_grad,
                    antrag.laengengrad_min,
                    antrag.laengengrad_sek,
                )

                sirene.save(
                    update_fields=[
                        "breitengrad",
                        "laengengrad",
                    ]
                )

                status, _ = SirenFRTStatus.objects.update_or_create(
                    sirene=sirene,
                    defaults={
                        "status": SirenFRTStatus.Status.BEANTRAGT,
                        "antragsdatum": antrag.antragsdatum,
                        "suchkreisname": antrag.suchkreisname,
                    },
                )

                apply_pending_freigabe(status)

                self.status_created += 1

        return self
    

def normalisiere_suchkreisname(name):
    if not name:
        return ""

    name = name.strip()

    if name.lower().startswith("sirene "):
        name = name[7:]

    return name.strip().lower()


def pending_siren():
    processed = 0

    for antrag in SirenFRTAntrag.objects.all():

        sirenen = finde_sirene(antrag)

        if len(sirenen) != 1:
            continue

        sirene = sirenen[0]

        status, _ = SirenFRTStatus.objects.update_or_create(
            sirene=sirene,
            defaults={
                "status": SirenFRTStatus.Status.BEANTRAGT,
                "antragsdatum": antrag.antragsdatum,
                "suchkreisname": antrag.suchkreisname,
            },
        )

        apply_pending_freigabe(status)
        
        antrag.delete()
        processed += 1

    return processed

class SirenFRTFreigabeImporter:

    def __init__(self, rows):
        self.rows = rows
        self.updated = 0
        self.not_found = 0

    def run(self):
        
        print(" SirenFRTFreigabeImporter WIRD AUFGERUFEN ")

        for row in self.rows:

            suchkreisname = row["suchkreisname"]
            
            print("FRT FREIGABE:", repr(suchkreisname))

            suchkreisname_freigabe = normalisiere_suchkreisname(
                suchkreisname
            )

            status = None

            for kandidat in SirenFRTStatus.objects.all():

                if (
                    normalisiere_suchkreisname(kandidat.suchkreisname)
                    == suchkreisname_freigabe
                ):
                    status = kandidat
                    break

            print("GEFUNDENER STATUS:", status)
            
            if not status:
                SirenFRTFreigabePending.objects.update_or_create(
                    suchkreisname=suchkreisname_freigabe,
                    defaults={
                        "freigabedatum": row["freigabedatum"],
                        "as_kommentar": row["as_kommentar"],
                        "zuteilungsnummer": row["zuteilungsnummer"],
                    },
                )
                self.not_found += 1
                continue

            status.status = SirenFRTStatus.Status.FREIGEGEBEN
            status.freigabedatum = row["freigabedatum"]
            status.as_kommentar = row["as_kommentar"]

            status.save(
                update_fields=[
                    "status",
                    "freigabedatum",
                    "as_kommentar",
                ]
            )
            
            self.updated += 1

        return self
    
    
def apply_pending_freigabe(status):
    suchkreisname = normalisiere_suchkreisname(
        status.suchkreisname
    )

    freigabe = SirenFRTFreigabePending.objects.filter(
        suchkreisname=suchkreisname
    ).first()

    if not freigabe:
        return False

    status.status = SirenFRTStatus.Status.FREIGEGEBEN
    status.freigabedatum = freigabe.freigabedatum
    status.as_kommentar = freigabe.as_kommentar
    status.save(
        update_fields=[
            "status",
            "freigabedatum",
            "as_kommentar",
        ]
    )

    freigabe.delete()

    return True